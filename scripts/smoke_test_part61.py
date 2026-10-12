#!/usr/bin/env python3
"""Part 61 smoke test — excel_write tool."""
import ast, base64, io, os, sys, asyncio
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
try:
    from dotenv import load_dotenv
    load_dotenv(ROOT / ".env")
except Exception:
    pass

PASS, FAIL = 0, 0


def check(label, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {label}")
    else:
        FAIL += 1
        print(f"  [FAIL] {label}" + (f" — {detail}" if detail else ""))


def main():
    print("=" * 62)
    print(" Part 61 — excel_write tool")
    print("=" * 62)

    tool_file = ROOT / "src/sunnyware/tools/excel_write.py"
    init_file = ROOT / "src/sunnyware/tools/__init__.py"

    check("excel_write.py exists", tool_file.exists())
    if not tool_file.exists():
        return 1

    src = tool_file.read_text(encoding="utf-8")
    init_src = init_file.read_text(encoding="utf-8")

    check("class defined", "class ExcelWriteTool" in src)
    check("name = excel_write", 'name = "excel_write"' in src)
    check("openpyxl Workbook", "Workbook" in src)
    check("styles import", "from openpyxl.styles import" in src)
    check("number format map", "_NUMFMT" in src)
    check("formula support", '"formula"' in src)
    check("workspace jail", "_workspace_root" in src)
    check("safe filename", "_safe_name" in src)
    check("sheet limits", "_MAX_SHEETS" in src and "_MAX_ROWS" in src)
    check("registered in __init__",
          "ExcelWriteTool.name: ExcelWriteTool()" in init_src)

    try:
        ast.parse(src)
        check("excel_write.py parses", True)
    except SyntaxError as e:
        check("excel_write.py parses", False, detail=str(e))

    try:
        import openpyxl  # noqa
        check("openpyxl importable", True)
    except Exception as e:
        check("openpyxl importable", False, detail=str(e))
        return 1 if FAIL else 0

    # Integration — build a 2-sheet workbook
    async def run(**kw):
        from sunnyware.tools.excel_write import ExcelWriteTool
        return await ExcelWriteTool().run(kw)

    r = asyncio.run(run(
        filename="smoke_p61.xlsx",
        sheets=[
            {
                "name": "P&L",
                "headers": ["Line", "FY24", "FY25E", "Variance %"],
                "rows": [
                    ["Revenue", 12000000, 15600000, {"formula": "=(C2-B2)/B2"}],
                    ["COGS", 7200000, 9100000, {"formula": "=(C3-B3)/B3"}],
                    ["Gross Profit", {"formula": "=B2-B3"},
                     {"formula": "=C2-C3"}, {"formula": "=(C4-B4)/B4"}],
                ],
                "column_widths": [22, 14, 14, 12],
                "number_formats": {"2": "money", "3": "money", "4": "percent"},
            },
            {
                "name": "Assumptions",
                "headers": ["Key", "Value"],
                "rows": [["Growth", 0.30], ["Churn", 0.05]],
                "number_formats": {"2": "percent"},
                "freeze": True,
            },
        ],
    ))
    check("workbook build succeeds", "error" not in r,
          detail=str(r)[:200] if r.get("error") else "")
    if "error" not in r:
        check("2 sheets written", r.get("sheets_written") == 2)
        check("5 rows written", r.get("rows_written") == 5,
              detail=f"got {r.get('rows_written')}")
        check("file on disk", Path(r["path"]).exists())
        check("size > 4 KB", r.get("size_bytes", 0) > 4000)
        check("returned base64", r.get("returned_base64") is True)

        # Verify by loading it back
        if r.get("base64"):
            try:
                from openpyxl import load_workbook
                raw = base64.b64decode(r["base64"])
                wb = load_workbook(io.BytesIO(raw))
                check("workbook loads back", True)
                check("sheet names preserved",
                      wb.sheetnames == ["P&L", "Assumptions"],
                      detail=str(wb.sheetnames))
                ws = wb["P&L"]
                check("header A1 == 'Line'", ws["A1"].value == "Line")
                check("header A1 bold", ws["A1"].font.bold is True)
                check("formula preserved in C4",
                      ws["C4"].value == "=C2-C3",
                      detail=repr(ws["C4"].value))
                check("freeze panes on sheet 1", ws.freeze_panes == "A2")
                check("number format applied to B2",
                      "money" in str(ws["B2"].number_format).lower()
                      or "#" in str(ws["B2"].number_format))
                wb.close()
            except Exception as e:
                check("workbook loads back", False, detail=str(e))

    # Rejections
    r2 = asyncio.run(run(sheets=[]))
    check("empty sheets rejected", "error" in r2)

    r3 = asyncio.run(run(sheets=[{"name": "x", "rows": []}],
                        filename="../escape.xlsx"))
    check("filename escape rejected OR sanitized",
          "error" in r3 or "/" not in r3.get("filename", ""))

    # Illegal sheet name should be sanitized, not error
    r4 = asyncio.run(run(sheets=[{"name": "a/b:c*?", "headers": ["h"],
                                  "rows": [[1]]}],
                        filename="p61_badname.xlsx"))
    check("illegal sheet name sanitized",
          "error" not in r4 and r4.get("sheets_written") == 1)

    # Cleanup
    ws = Path(os.getenv("SUNNYWARE_WORKSPACE") or "/data/workspace").resolve()
    for n in ("smoke_p61.xlsx", "p61_badname.xlsx"):
        try: (ws / n).unlink()
        except Exception: pass

    print()
    print("=" * 62)
    print(f" RESULT: {PASS} passed, {FAIL} failed")
    print("=" * 62)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
