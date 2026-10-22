# SPDX-FileCopyrightText: 2026 Shamiur Rashid Sunny
# SPDX-License-Identifier: AGPL-3.0-only
"""Excel reader tool (Part 62).

Read .xlsx workbooks from the workspace. Returns sheet names, headers,
row data as JSON, and optionally raw formulas. Caps rows/cols per sheet
to keep responses small. openpyxl is preinstalled — zero new deps.
"""
import os
from pathlib import Path
from typing import Optional

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except Exception:
    _HAS_OPENPYXL = False


_MAX_ROWS_PER_SHEET = 2000
_MAX_COLS = 100
_MAX_SHEETS = 20
_MAX_CELL_CHARS = 500


def _workspace_root() -> Path:
    p = os.getenv("SUNNYWARE_WORKSPACE") or os.getenv("WORKSPACE_PATH")
    if p:
        return Path(p).resolve()
    return Path("/data/workspace").resolve()


def _coerce(v):
    """Return JSON-safe representation of a cell value."""
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    try:
        # datetime, date, time
        return v.isoformat()
    except Exception:
        s = str(v)
        return s[:_MAX_CELL_CHARS] if len(s) > _MAX_CELL_CHARS else s


def _sheet_to_records(ws, max_rows: int, include_formulas: bool) -> dict:
    """Return {headers, rows, truncated, cells} for a worksheet."""
    # Find used range explicitly (ws.dimensions can be wrong for sparse sheets)
    max_col = 0
    max_row = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                if cell.row > max_row:
                    max_row = cell.row
                if cell.column > max_col:
                    max_col = cell.column
    if max_col == 0 or max_row == 0:
        return {"headers": [], "rows": [], "truncated": False, "cells": 0}

    max_col = min(max_col, _MAX_COLS)
    truncated = max_row - 1 > max_rows
    rows_to_read = min(max_row, max_rows + 1)

    # Headers = row 1
    header_cells = [ws.cell(row=1, column=c) for c in range(1, max_col + 1)]
    headers = []
    for c_idx, cell in enumerate(header_cells, start=1):
        hv = cell.value
        if hv is None:
            headers.append(get_column_letter(c_idx))
        else:
            headers.append(str(hv))

    rows = []
    for r in range(2, rows_to_read + 1):
        record = {}
        for c_idx in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c_idx)
            v = cell.value
            # If formula mode, detect formula strings (openpyxl stores them
            # as '=...' when data_only=False)
            if isinstance(v, str) and v.startswith("=") and not include_formulas:
                record[headers[c_idx - 1]] = {
                    "formula": v,
                    "value": None,
                }
            else:
                record[headers[c_idx - 1]] = _coerce(v)
        rows.append(record)

    return {
        "headers": headers,
        "rows": rows,
        "truncated": truncated,
        "cells": (rows_to_read - 1) * max_col,
    }


class ExcelReadTool:
    name = "excel_read"
    description = "Excel spreadsheet reader for .xlsx files."
    parameters = {
        "file": {
            "type": "string",
            "description": "Workspace .xlsx filename",
            "required": True,
        },
        "sheet": {
            "type": "string",
            "description": "Specific sheet name (default: all sheets, up to 20)",
            "required": False,
        },
        "max_rows": {
            "type": "integer",
            "description": "Max rows per sheet (default 200, max 2000)",
            "required": False,
        },
        "include_formulas": {
            "type": "boolean",
            "description": "Keep formula strings instead of wrapping them (default false)",
            "required": False,
        },
        "list_only": {
            "type": "boolean",
            "description": "Only list sheet names + dimensions, skip row data (default false)",
            "required": False,
        },
    }

    async def run(self, args: dict) -> dict:
        if not _HAS_OPENPYXL:
            return {"error": "openpyxl not installed in this environment"}

        fname = (args.get("file") or "").strip()
        if not fname:
            return {"error": "file is required"}
        if not fname.lower().endswith(".xlsx"):
            return {"error": "only .xlsx files are supported"}

        root = _workspace_root()
        candidate = (root / fname).resolve()
        try:
            candidate.relative_to(root)
        except ValueError:
            return {"error": "file must live inside the workspace"}
        if not candidate.exists():
            return {"error": f"file not found: {fname}"}

        try:
            max_rows = int(args.get("max_rows") or 200)
        except (TypeError, ValueError):
            return {"error": "max_rows must be an integer"}
        max_rows = max(1, min(max_rows, _MAX_ROWS_PER_SHEET))

        include_formulas = bool(args.get("include_formulas", False))
        list_only = bool(args.get("list_only", False))
        wanted_sheet = (args.get("sheet") or "").strip()

        try:
            # read_only=True drops formulas in some xlsx — use normal mode.
            # data_only=False preserves formulas (as '=...' strings).
            wb = load_workbook(str(candidate), data_only=False,
                               read_only=False, keep_links=False)
        except Exception as e:
            return {"error": f"open failed: {type(e).__name__}: {e}"}

        sheet_names = wb.sheetnames
        if wanted_sheet:
            if wanted_sheet not in sheet_names:
                wb.close()
                return {"error": f"sheet not found: {wanted_sheet}",
                        "available_sheets": sheet_names}

        sheets_out = []
        targets = [wanted_sheet] if wanted_sheet else sheet_names[:_MAX_SHEETS]

        try:
            for name in targets:
                ws = wb[name]
                entry = {
                    "name": name,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                }
                if not list_only:
                    data = _sheet_to_records(ws, max_rows, include_formulas)
                    entry.update(data)
                sheets_out.append(entry)
        except Exception as e:
            wb.close()
            return {"error": f"read failed: {type(e).__name__}: {e}"}

        try:
            wb.close()
        except Exception:
            pass

        try:
            size = candidate.stat().st_size
        except Exception:
            size = 0

        return {
            "file": fname,
            "size_bytes": size,
            "sheet_count": len(sheet_names),
            "sheet_names": sheet_names,
            "sheets": sheets_out,
            "list_only": list_only,
        }
